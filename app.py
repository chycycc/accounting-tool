import json
import os
import re
import uuid
from datetime import datetime
from io import BytesIO

from flask import Flask, jsonify, request, render_template, send_file
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

app = Flask(__name__)

DATA_FILE = os.path.join(os.path.dirname(__file__), 'data.json')


def load_data():
    if os.path.exists(DATA_FILE):
        with open(DATA_FILE, 'r', encoding='utf-8') as f:
            return json.load(f)
    return {"prices": [150, 180, 300], "records": {}}


def save_data(data):
    with open(DATA_FILE, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=2)


def parse_raw_input(raw_str):
    """Parse input like '750*5+422' or '670 800 670 670' or '183包' or '183b' into kg total."""
    raw_str = raw_str.strip().replace('×', '*').replace('x', '*').replace('X', '*')
    nums = []
    bag_kg = 0
    for part in re.split(r'[+]', raw_str):
        part = part.strip()
        if not part:
            continue
        # Handle 包 unit: 183包 or 183b => kg = 183/40*1000
        bag_match = re.match(r'^(\d+)\s*(包|b)$', part, re.IGNORECASE)
        if bag_match:
            bag_kg += round(int(bag_match.group(1)) / 40 * 1000)
            continue
        if '*' in part:
            nums_part = part.split('*')
            if len(nums_part) == 2:
                try:
                    val = int(nums_part[0].strip())
                    count = int(nums_part[1].strip())
                    nums.extend([val] * count)
                except ValueError:
                    pass
        else:
            for num in re.split(r'[\s,，]+', part):
                num = num.strip()
                if num:
                    try:
                        nums.append(int(num))
                    except ValueError:
                        pass
    return sum(nums) + bag_kg


@app.route('/')
def index():
    return render_template('index.html')


@app.route('/api/prices', methods=['GET'])
def get_prices():
    data = load_data()
    return jsonify(data['prices'])


@app.route('/api/prices', methods=['POST'])
def add_price():
    data = load_data()
    price = request.json.get('price')
    if price and price not in data['prices']:
        data['prices'].append(price)
        data['prices'].sort()
        save_data(data)
    return jsonify(data['prices'])


@app.route('/api/prices/<int:price>', methods=['DELETE'])
def delete_price(price):
    data = load_data()
    if price in data['prices']:
        data['prices'].remove(price)
        save_data(data)
    return jsonify(data['prices'])


@app.route('/api/months', methods=['GET'])
def get_months():
    data = load_data()
    months = sorted(data['records'].keys(), reverse=True)
    return jsonify(months)


def date_sort_key(r):
    """按日期排序，如 7.1, 7.2, 7.10, 7.12"""
    parts = r['date'].split('.')
    return (int(parts[0]), int(parts[1])) if len(parts) == 2 else (0, 0)


@app.route('/api/records', methods=['GET'])
def get_records():
    month = request.args.get('month')
    data = load_data()
    records = data['records'].get(month, [])
    records.sort(key=date_sort_key)
    return jsonify(records)


@app.route('/api/records', methods=['POST'])
def add_record():
    data = load_data()
    req = request.json
    month = req.get('month')
    date = req.get('date')
    raw = req.get('raw')
    price = int(req.get('price'))

    kg = parse_raw_input(raw)

    record = {
        'id': str(uuid.uuid4())[:8],
        'date': date,
        'raw': raw,
        'kg': kg,
        'price': price
    }

    if month not in data['records']:
        data['records'][month] = []
    data['records'][month].append(record)
    save_data(data)
    return jsonify(record), 201


@app.route('/api/records/<record_id>', methods=['DELETE'])
def delete_record(record_id):
    data = load_data()
    for month, records in data['records'].items():
        for i, r in enumerate(records):
            if r['id'] == record_id:
                records.pop(i)
                save_data(data)
                return jsonify({'ok': True})
    return jsonify({'error': 'not found'}), 404


@app.route('/api/report', methods=['GET'])
def get_report():
    year = request.args.get('year')
    data = load_data()
    report = {'months': [], 'by_price': {}}

    for month_key in sorted(data['records'].keys()):
        if year and not month_key.startswith(year):
            continue
        records = data['records'][month_key]
        total_kg = sum(r['kg'] for r in records)
        total_tons = total_kg / 1000
        total_amount = sum(r['kg'] / 1000 * r['price'] for r in records)

        price_breakdown = {}
        for r in records:
            p = r['price']
            if p not in price_breakdown:
                price_breakdown[p] = {'kg': 0, 'amount': 0}
            price_breakdown[p]['kg'] += r['kg']
            price_breakdown[p]['amount'] += r['kg'] / 1000 * p

        report['months'].append({
            'month': month_key,
            'days': len(records),
            'total_kg': total_kg,
            'total_tons': round(total_tons, 3),
            'total_amount': round(total_amount, 2),
            'by_price': {str(k): v for k, v in price_breakdown.items()}
        })

        for p, v in price_breakdown.items():
            if p not in report['by_price']:
                report['by_price'][p] = {'kg': 0, 'amount': 0}
            report['by_price'][p]['kg'] += v['kg']
            report['by_price'][p]['amount'] += v['amount']

    for p in report['by_price']:
        report['by_price'][p]['amount'] = round(report['by_price'][p]['amount'], 2)

    return jsonify(report)


@app.route('/api/top10', methods=['GET'])
def get_top10():
    year = request.args.get('year')
    data = load_data()
    all_records = []
    for month_key, records in data['records'].items():
        if year and not month_key.startswith(year):
            continue
        for r in records:
            all_records.append({**r, 'month': month_key})
    all_records.sort(key=lambda x: x['kg'], reverse=True)
    return jsonify(all_records[:10])


@app.route('/api/export', methods=['GET'])
def export_excel():
    month = request.args.get('month')
    data = load_data()
    records = sorted(data['records'].get(month, []), key=date_sort_key)

    wb = Workbook()
    ws = wb.active
    ws.title = f'{month}账目'

    header_font = Font(bold=True, size=12)
    blue_font = Font(size=11, color='0000FF')
    black_font = Font(size=11, color='000000')
    thin = Border(left=Side('thin'), right=Side('thin'), top=Side('thin'), bottom=Side('thin'))
    header_fill = PatternFill('solid', fgColor='D9E1F2')
    sum_fill = PatternFill('solid', fgColor='FFF2CC')
    center = Alignment(horizontal='center', vertical='center')

    headers = ['号期', '原始数据', '当号合计(kg)', '吨数', '单价(元/吨)', '金额(元)']
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = header_font
        c.fill = header_fill
        c.border = thin
        c.alignment = center

    widths = [10, 30, 16, 12, 14, 14]
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w

    for i, r in enumerate(records, 2):
        ws.cell(row=i, column=1, value=r['date']).border = thin
        ws.cell(row=i, column=1).alignment = center
        ws.cell(row=i, column=2, value=r['raw']).border = thin
        ws.cell(row=i, column=3, value=r['kg']).font = blue_font
        ws.cell(row=i, column=3).border = thin
        ws.cell(row=i, column=3).alignment = center
        ws.cell(row=i, column=3).number_format = '#,##0'
        ws.cell(row=i, column=4, value=f'=C{i}/1000').font = black_font
        ws.cell(row=i, column=4).border = thin
        ws.cell(row=i, column=4).alignment = center
        ws.cell(row=i, column=4).number_format = '0.000'
        ws.cell(row=i, column=5, value=r['price']).font = blue_font
        ws.cell(row=i, column=5).border = thin
        ws.cell(row=i, column=5).alignment = center
        ws.cell(row=i, column=5).number_format = '#,##0'
        ws.cell(row=i, column=6, value=f'=D{i}*E{i}').font = black_font
        ws.cell(row=i, column=6).border = thin
        ws.cell(row=i, column=6).alignment = center
        ws.cell(row=i, column=6).number_format = '#,##0.00'

    lr = len(records) + 1
    sr = lr + 2
    labels = ['合计', '总天数', '总kg', '总吨', '总金额']
    for j, label in enumerate(labels):
        ws.merge_cells(start_row=sr+j, start_column=1, end_row=sr+j, end_column=2)
        for c in range(1, 3):
            ws.cell(row=sr+j, column=c).fill = sum_fill
            ws.cell(row=sr+j, column=c).border = thin
        ws.cell(row=sr+j, column=1, value=label).font = Font(bold=True, size=11)
        ws.cell(row=sr+j, column=1).alignment = center

    ws.cell(row=sr+1, column=3, value=f'=COUNTA(A2:A{lr})').font = black_font
    ws.cell(row=sr+1, column=3).fill = sum_fill
    ws.cell(row=sr+1, column=3).border = thin
    ws.cell(row=sr+1, column=3).alignment = center

    ws.cell(row=sr+2, column=3, value=f'=SUM(C2:C{lr})').font = black_font
    ws.cell(row=sr+2, column=3).fill = sum_fill
    ws.cell(row=sr+2, column=3).border = thin
    ws.cell(row=sr+2, column=3).alignment = center
    ws.cell(row=sr+2, column=3).number_format = '#,##0'

    ws.cell(row=sr+3, column=4, value=f'=SUM(D2:D{lr})').font = black_font
    ws.cell(row=sr+3, column=4).fill = sum_fill
    ws.cell(row=sr+3, column=4).border = thin
    ws.cell(row=sr+3, column=4).alignment = center
    ws.cell(row=sr+3, column=4).number_format = '0.000'

    ws.cell(row=sr+4, column=6, value=f'=SUM(F2:F{lr})').font = Font(bold=True, size=12, color='FF0000')
    ws.cell(row=sr+4, column=6).fill = sum_fill
    ws.cell(row=sr+4, column=6).border = thin
    ws.cell(row=sr+4, column=6).alignment = center
    ws.cell(row=sr+4, column=6).number_format = '#,##0.00'

    for j in range(5):
        for c in [3, 4, 5, 6]:
            ws.cell(row=sr+j, column=c).fill = sum_fill
            ws.cell(row=sr+j, column=c).border = thin

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True, download_name=f'{month}账目.xlsx')


@app.route('/api/export_report', methods=['GET'])
def export_report_excel():
    year = request.args.get('year')
    data = load_data()
    wb = Workbook()
    ws = wb.active
    ws.title = f'{year}年报表'

    header_font = Font(bold=True, size=12)
    black_font = Font(size=11, color='000000')
    thin = Border(left=Side('thin'), right=Side('thin'), top=Side('thin'), bottom=Side('thin'))
    header_fill = PatternFill('solid', fgColor='D9E1F2')
    sum_fill = PatternFill('solid', fgColor='FFF2CC')
    center = Alignment(horizontal='center', vertical='center')

    headers = ['月份', '天数', '总kg', '总吨', '总金额(元)']
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = header_font
        c.fill = header_fill
        c.border = thin
        c.alignment = center

    widths = [12, 10, 14, 12, 16]
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w

    row_idx = 2
    months_data = []
    for month_key in sorted(data['records'].keys()):
        if year and not month_key.startswith(year):
            continue
        records = data['records'][month_key]
        total_kg = sum(r['kg'] for r in records)
        total_tons = total_kg / 1000
        total_amount = sum(r['kg'] / 1000 * r['price'] for r in records)
        months_data.append((month_key, len(records), total_kg, round(total_tons, 3), round(total_amount, 2)))

    for md in months_data:
        for col, val in enumerate(md, 1):
            c = ws.cell(row=row_idx, column=col, value=val)
            c.font = black_font
            c.border = thin
            c.alignment = center
        ws.cell(row=row_idx, column=3).number_format = '#,##0'
        ws.cell(row=row_idx, column=4).number_format = '0.000'
        ws.cell(row=row_idx, column=5).number_format = '#,##0.00'
        row_idx += 1

    sr = row_idx + 1
    ws.cell(row=sr, column=1, value='全年合计').font = Font(bold=True, size=11)
    ws.cell(row=sr, column=1).fill = sum_fill
    ws.cell(row=sr, column=1).border = thin
    ws.cell(row=sr, column=1).alignment = center
    ws.cell(row=sr, column=2, value=f'=SUM(B2:B{row_idx-1})').fill = sum_fill
    ws.cell(row=sr, column=2).border = thin
    ws.cell(row=sr, column=2).alignment = center
    ws.cell(row=sr, column=3, value=f'=SUM(C2:C{row_idx-1})').fill = sum_fill
    ws.cell(row=sr, column=3).border = thin
    ws.cell(row=sr, column=3).number_format = '#,##0'
    ws.cell(row=sr, column=3).alignment = center
    ws.cell(row=sr, column=4, value=f'=SUM(D2:D{row_idx-1})').fill = sum_fill
    ws.cell(row=sr, column=4).border = thin
    ws.cell(row=sr, column=4).number_format = '0.000'
    ws.cell(row=sr, column=4).alignment = center
    ws.cell(row=sr, column=5, value=f'=SUM(E2:E{row_idx-1})').fill = sum_fill
    ws.cell(row=sr, column=5).border = thin
    ws.cell(row=sr, column=5).number_format = '#,##0.00'
    ws.cell(row=sr, column=5).alignment = center
    ws.cell(row=sr, column=5).font = Font(bold=True, size=12, color='FF0000')

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True, download_name=f'{year}年报表.xlsx')


if __name__ == '__main__':
    app.run(host='0.0.0.0', port=5000, debug=True)
